import os
import uuid
import logging
from datetime import datetime

import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm

logger = logging.getLogger(__name__)


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _export_path(folder: str, prefix: str, ext: str) -> str:
    """Generate a unique export file path."""
    os.makedirs(folder, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = f"{prefix}_{ts}_{uuid.uuid4().hex[:6]}.{ext}"
    return os.path.join(folder, name)


# ─── Excel Export ─────────────────────────────────────────────────────────────

def export_excel(df: pd.DataFrame, metrics: dict, behavior: dict, output_folder: str = "exports") -> str:
    if df.empty:
        raise ValueError("Cannot export empty dataframe")
    path = _export_path(output_folder, "tradeflow", "xlsx")

    GREEN = "FF26A69A"
    RED   = "FFEF5350"
    DARK  = "FF1E1E2E"
    HEADER_COLOR = "FF3949AB"

    with pd.ExcelWriter(path, engine="openpyxl") as writer:

        # ── Sheet 1: Trades ──────────────────────────────────────────────────
        df.to_excel(writer, sheet_name="Trades", index=False)
        ws = writer.sheets["Trades"]
        _style_header_row(ws, HEADER_COLOR)
        _color_pnl_column(ws, df, GREEN, RED)
        _autofit_columns(ws)

        # ── Sheet 2: Performance Metrics ─────────────────────────────────────
        metrics_df = pd.DataFrame([
            {"Metric": k.replace("_", " ").title(), "Value": v}
            for k, v in metrics.items()
        ])
        metrics_df.to_excel(writer, sheet_name="Performance", index=False)
        ws2 = writer.sheets["Performance"]
        _style_header_row(ws2, HEADER_COLOR)
        _autofit_columns(ws2)

        # ── Sheet 3: Behavior ─────────────────────────────────────────────────
        behavior_rows = []
        for k, v in behavior.items():
            if isinstance(v, dict):
                continue   # skip nested dicts (weekday_pnl)
            behavior_rows.append(
                {"Metric": k.replace("_", " ").title(), "Value": str(v)}
            )
        beh_df = pd.DataFrame(behavior_rows)
        beh_df.to_excel(writer, sheet_name="Behavior", index=False)
        ws3 = writer.sheets["Behavior"]
        _style_header_row(ws3, HEADER_COLOR)
        _autofit_columns(ws3)

    logger.info(f"Excel exported: {path}")
    return path


def _style_header_row(ws, color_hex: str):
    fill   = PatternFill("solid", fgColor=color_hex)
    font   = Font(bold=True, color="FFFFFFFF")
    align  = Alignment(horizontal="center")
    for cell in ws[1]:
        cell.fill  = fill
        cell.font  = font
        cell.alignment = align


def _color_pnl_column(ws, df: pd.DataFrame, green: str, red: str):
    if "pnl" not in df.columns:
        return
    pnl_idx = df.columns.tolist().index("pnl") + 1   # 1-indexed
    green_fill = PatternFill("solid", fgColor=green)
    red_fill   = PatternFill("solid", fgColor=red)
    for row in ws.iter_rows(min_row=2, min_col=pnl_idx, max_col=pnl_idx):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.fill = green_fill if cell.value >= 0 else red_fill


def _autofit_columns(ws, min_width=10, max_width=40):
    for col in ws.columns:
        length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = (
            max(min_width, min(length + 2, max_width))
        )


# ─── PDF Export ───────────────────────────────────────────────────────────────

def export_pdf(metrics: dict, behavior: dict, insights: list, output_folder: str = "exports") -> str:
    path = _export_path(output_folder, "tradeflow", "pdf")

    doc = SimpleDocTemplate(
        path, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=2*cm, bottomMargin=2*cm
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "TFTitle", parent=styles["Title"],
        textColor=colors.HexColor("#5c6bc0"),
        fontSize=22, spaceAfter=6
    )
    section_style = ParagraphStyle(
        "TFSection", parent=styles["Heading2"],
        textColor=colors.HexColor("#3949ab"),
        spaceBefore=14, spaceAfter=6
    )
    body_style = styles["BodyText"]

    insight_colors = {
        "danger":  colors.HexColor("#ef5350"),
        "warning": colors.HexColor("#ffa726"),
        "info":    colors.HexColor("#26a69a"),
    }

    story = []

    # ── Header ────────────────────────────────────────────────────────────────
    story.append(Paragraph("TradeFlow", title_style))
    story.append(Paragraph(
        f"Performance Report — Generated {datetime.now().strftime('%d %b %Y %H:%M')}",
        styles["Normal"]
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#3949ab")))
    story.append(Spacer(1, 14))

    # ── Performance Metrics Table ─────────────────────────────────────────────
    story.append(Paragraph("Performance Metrics", section_style))
    metrics_data = [["Metric", "Value"]] + [
        [k.replace("_", " ").title(), str(v)]
        for k, v in metrics.items()
    ]
    story.append(_build_table(metrics_data))
    story.append(Spacer(1, 14))

    # ── Behavioral Analytics Table ────────────────────────────────────────────
    story.append(Paragraph("Behavioral Analytics", section_style))
    behavior_data = [["Metric", "Value"]] + [
        [k.replace("_", " ").title(), str(v)]
        for k, v in behavior.items()
        if not isinstance(v, dict)
    ]
    story.append(_build_table(behavior_data))
    story.append(Spacer(1, 14))

    # ── AI Insights ───────────────────────────────────────────────────────────
    story.append(Paragraph("AI Insights", section_style))
    if not insights:
        insights = [{"text": "No AI insights available.", "level": "info"}]
    for item in insights:
        # item is {"text": str, "level": str}
        if isinstance(item, dict):
            text = item.get("text", str(item))
            level = item.get("level", "info")
        else:
            text = str(item)
            level = "info"
        level = item.get("level", "info")
        color = insight_colors.get(level, colors.black)
        style = ParagraphStyle(
            f"ins_{level}",
            parent=body_style,
            textColor=color,
            leftIndent=10,
            spaceAfter=4
        )
        prefix = {"danger": "⚠ ", "warning": "→ ", "info": "✓ "}.get(level, "• ")
        story.append(Paragraph(f"{prefix}{text}", style))

    doc.build(story)
    logger.info(f"PDF exported: {path}")
    return path


def _build_table(data: list) -> Table:
    t = Table(data, colWidths=[9*cm, 7*cm])
    style = TableStyle([
        ("BACKGROUND",  (0,0), (-1,0),  colors.HexColor("#3949ab")),
        ("TEXTCOLOR",   (0,0), (-1,0),  colors.white),
        ("FONTNAME",    (0,0), (-1,0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0,0), (-1,-1), 9),
        ("ROWBACKGROUNDS", (0,1), (-1,-1),
         [colors.HexColor("#1e1e2e"), colors.HexColor("#252535")]),
        ("TEXTCOLOR",   (0,1), (-1,-1), colors.HexColor("#e0e0e0")),
        ("GRID",        (0,0), (-1,-1), 0.25, colors.HexColor("#3949ab")),
        ("ALIGN",       (1,0), (1,-1),  "RIGHT"),
        ("VALIGN",      (0,0), (-1,-1), "MIDDLE"),
        ("ROWHEIGHT",   (0,0), (-1,-1), 18),
    ])
    t.setStyle(style)
    return t